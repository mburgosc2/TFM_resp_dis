"""Genera el Excel reproducible de experimentos Stage 1 (no tos / tos)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent
OUTPUT_PATH = ROOT / "stage1_cough_no_cough_experiments_summary.xlsx"
HISTORICAL_TEXT = ROOT / "graphs_results_experiment_random" / "results_training.txt"
CURRENT_RESULTS = ROOT / "results_stage1_cough_no_cough" / "mfcc117_rf" / "full"
HISTORICAL_QUALITY_TEXT = ROOT / "graphs_results_experiment_poor" / "results_training.txt"
QUALITY_RESULTS = (
    ROOT
    / "results_stage1_cough_no_cough"
    / "mfcc117_rf_audio_quality_poor"
    / "full"
)
COMPACT_LR_RESULTS = (
    ROOT
    / "results_stage1_cough_no_cough"
    / "mfcc117_lr_audio_quality_poor"
    / "full"
)
COMPACT_RF_RESULTS = (
    ROOT
    / "results_stage1_cough_no_cough"
    / "mfcc117_rf_compact_audio_quality_poor"
    / "full"
)
RANDOM_LR_RESULTS = (
    ROOT / "results_stage1_cough_no_cough" / "mfcc117_lr_random" / "full"
)
RANDOM_RF_RESULTS = (
    ROOT / "results_stage1_cough_no_cough" / "mfcc117_rf_compact_random" / "full"
)
FSD50K_COUGH_LR_RESULTS = (
    ROOT
    / "results_stage1_cough_no_cough"
    / "mfcc117_lr_fsd50k_coughs_augmented_random"
    / "full"
)
FSD50K_COUGH_LR_SEARCH_RESULTS = (
    ROOT
    / "results_stage1_cough_no_cough"
    / "mfcc117_lr_fsd50k_coughs_search_random"
    / "full"
)
WST_O01_RECORDING_LR_RESULTS = (
    ROOT
    / "results_stage1_cough_no_cough"
    / "wst_o01_recording_enhanced_lr_fsd50k_coughs_random"
    / "full"
)


def historical_metrics() -> tuple[dict[str, Any], pd.DataFrame, pd.DataFrame]:
    summary: dict[str, Any] = {
        "ID": "S1-H01",
        "Estado": "Histórico",
        "Experimento": "MFCC117 + RF (splits antiguos)",
        "Features": "13 MFCC + delta + delta2; mean/std/max = 117",
        "Clasificador": "Random Forest",
        "Datos/splits": "features_extracted_experiment_random (10-08-2026)",
        "Umbral": 0.5,
        "Nota": "Resultado original; metadatos anteriores al consenso actualizado.",
        "Fuente": str(HISTORICAL_TEXT),
    }
    details: list[dict[str, Any]] = []
    folds: list[dict[str, Any]] = []
    if not HISTORICAL_TEXT.is_file():
        summary["Estado"] = "No encontrado"
        return summary, pd.DataFrame(), pd.DataFrame()

    text = HISTORICAL_TEXT.read_text(encoding="utf-8", errors="ignore")
    for fold, score in re.findall(r"Fold\s+(\d+)\s+-\s+F1-Score:\s*([0-9.]+)", text):
        folds.append({"ID": "S1-H01", "fold": int(fold), "f1_cough": float(score)})
    match = re.search(r"VALIDACI[^:]*:\s*([0-9.]+)", text, flags=re.IGNORECASE)
    if match:
        summary["F1 CV medio"] = float(match.group(1))
    if folds:
        fold_scores = [row["f1_cough"] for row in folds]
        summary["F1 CV std"] = float(np.std(fold_scores, ddof=0))

    block_pattern = re.compile(
        r"Conjunto\s+(TRAIN|VALIDATION|TEST \(BLIND\))\s*:(.*?)(?=Conjunto|Pipeline|\Z)",
        flags=re.DOTALL,
    )
    metric_patterns = {
        "accuracy": r"Accuracy:\s*([0-9.]+)",
        "precision_cough": r"Precision:\s*([0-9.]+)",
        "recall_cough": r"Recall:\s*([0-9.]+)",
        "f1_cough": r"F1-Score:\s*([0-9.]+)",
        "roc_auc": r"ROC-AUC:\s*([0-9.]+)",
    }
    split_names = {"TRAIN": "train_fit", "VALIDATION": "validation", "TEST (BLIND)": "test"}
    for raw_split, block in block_pattern.findall(text):
        row: dict[str, Any] = {"ID": "S1-H01", "dataset": split_names[raw_split]}
        for name, pattern in metric_patterns.items():
            found = re.search(pattern, block)
            row[name] = float(found.group(1)) if found else np.nan
        details.append(row)

    detail_df = pd.DataFrame(details)
    validation = detail_df[detail_df["dataset"] == "validation"]
    test = detail_df[detail_df["dataset"] == "test"]
    if not validation.empty:
        row = validation.iloc[0]
        summary.update(
            {
                "F1 validation": row.get("f1_cough", np.nan),
                "Recall tos validation": row.get("recall_cough", np.nan),
                "Precision tos validation": row.get("precision_cough", np.nan),
                "AUC validation": row.get("roc_auc", np.nan),
            }
        )
    if not test.empty:
        row = test.iloc[0]
        summary.update(
            {
                "F1 test": row.get("f1_cough", np.nan),
                "Recall tos test": row.get("recall_cough", np.nan),
                "AUC test": row.get("roc_auc", np.nan),
            }
        )
    return summary, detail_df, pd.DataFrame(folds)


def current_metrics() -> tuple[dict[str, Any], pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    metrics_path = CURRENT_RESULTS / "metrics_summary.csv"
    summary: dict[str, Any] = {
        "ID": "S1-C01",
        "Estado": "Pendiente de ejecución",
        "Experimento": "MFCC117 + RF (splits actuales)",
        "Features": "13 MFCC + delta + delta2; mean/std/max = 117",
        "Clasificador": "Random Forest",
        "Datos/splits": "multiclass_4c actual; 0=no_tos, 1/2/3=tos",
        "Umbral": 0.5,
        "Nota": "Repetición controlada con consenso y folds actuales.",
        "Fuente": str(metrics_path),
    }
    if not metrics_path.is_file():
        return summary, pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    metrics = pd.read_csv(metrics_path)
    metrics.insert(0, "ID", "S1-C01")
    folds_path = CURRENT_RESULTS / "cv_fold_metrics.csv"
    folds = pd.read_csv(folds_path) if folds_path.is_file() else pd.DataFrame()
    if not folds.empty:
        folds.insert(0, "ID", "S1-C01")
    config_path = CURRENT_RESULTS / "experiment_configuration.csv"
    config = pd.read_csv(config_path) if config_path.is_file() else pd.DataFrame()
    if not config.empty:
        config.insert(0, "ID", "S1-C01")

    summary["Estado"] = "Finalizado"
    oof = metrics[metrics["dataset"] == "train_oof"]
    validation = metrics[metrics["dataset"] == "validation"]
    test = metrics[metrics["dataset"] == "test"]
    if not oof.empty:
        summary["Macro-F1 OOF"] = oof.iloc[0].get("macro_f1", np.nan)
    if not folds.empty:
        summary["F1 CV medio"] = folds["f1_cough"].mean()
        summary["F1 CV std"] = folds["f1_cough"].std(ddof=0)
    if not validation.empty:
        row = validation.iloc[0]
        summary.update(
            {
                "F1 validation": row.get("f1_cough", np.nan),
                "Macro-F1 validation": row.get("macro_f1", np.nan),
                "Recall tos validation": row.get("recall_cough", np.nan),
                "Especificidad validation": row.get("specificity_no_cough", np.nan),
                "Precision tos validation": row.get("precision_cough", np.nan),
                "AUC validation": row.get("roc_auc", np.nan),
            }
        )
    if not test.empty:
        row = test.iloc[0]
        summary.update(
            {
                "F1 test": row.get("f1_cough", np.nan),
                "Macro-F1 test": row.get("macro_f1", np.nan),
                "Recall tos test": row.get("recall_cough", np.nan),
                "Especificidad test": row.get("specificity_no_cough", np.nan),
                "Precision tos test": row.get("precision_cough", np.nan),
                "AUC test": row.get("roc_auc", np.nan),
            }
        )
    if not config.empty:
        config_map = dict(zip(config["parameter"], config["value"]))
        summary["Configuracion ganadora"] = (
            f"RF trees={config_map.get('n_estimators', '?')}, "
            f"depth={config_map.get('max_depth', '?')}, "
            f"leaf={config_map.get('min_samples_leaf', '?')}"
        )
        if "model_size_kb" in config_map:
            summary["Tamano modelo KB"] = config_map["model_size_kb"]
        if "tree_nodes_total" in config_map:
            summary["Nodos arboles"] = config_map["tree_nodes_total"]
    return summary, metrics, folds, config


def add_table(ws, name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    table = Table(displayName=name, ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)


def style_workbook() -> None:
    workbook = load_workbook(OUTPUT_PATH)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in ws.columns:
            values = [str(cell.value or "") for cell in column[: min(ws.max_row, 100)]]
            width = min(max(max(map(len, values), default=8) + 2, 10), 45)
            ws.column_dimensions[get_column_letter(column[0].column)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    summary = workbook["Resumen"]
    header = {cell.value: cell.column for cell in summary[1]}
    for metric in ["F1 validation", "Recall tos validation", "AUC validation"]:
        if metric in header and summary.max_row >= 3:
            letter = get_column_letter(header[metric])
            summary.conditional_formatting.add(
                f"{letter}2:{letter}{summary.max_row}",
                ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile",
                               mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
            )
    if "F1 validation" in header:
        chart = BarChart()
        chart.title = "F1 de tos en validation"
        chart.y_axis.title = "F1"
        chart.height = 7
        chart.width = 14
        chart.add_data(
            Reference(summary, min_col=header["F1 validation"], min_row=1, max_row=summary.max_row),
            titles_from_data=True,
        )
        chart.set_categories(Reference(summary, min_col=3, min_row=2, max_row=summary.max_row))
        summary.add_chart(chart, f"A{summary.max_row + 4}")
    workbook.save(OUTPUT_PATH)


def build_workbook() -> Path:
    historical_summary, historical_detail, historical_folds = historical_metrics()
    current_summary, current_detail, current_folds, current_config = current_metrics()
    summary = pd.DataFrame([historical_summary, current_summary])
    preferred_columns = [
        "ID", "Estado", "Experimento", "Features", "Clasificador", "Datos/splits", "Umbral",
        "F1 CV medio", "F1 CV std", "Macro-F1 OOF", "F1 validation", "Macro-F1 validation",
        "Recall tos validation", "Especificidad validation", "Precision tos validation",
        "AUC validation", "F1 test", "Recall tos test", "AUC test", "Nota", "Fuente",
    ]
    summary = summary.reindex(columns=preferred_columns)
    details = pd.concat([historical_detail, current_detail], ignore_index=True, sort=False)
    folds = pd.concat([historical_folds, current_folds], ignore_index=True, sort=False)
    notes = pd.DataFrame(
        {
            "Aspecto": ["Objetivo", "Mapeo", "Comparabilidad", "TEST", "Actualización"],
            "Descripción": [
                "Stage 1 detecta tos frente a no tos.",
                "Clase 0=no_cough; dry, wet y unknown se consideran tos (clase 1).",
                "S1-H01 usa splits antiguos; S1-C01 usa splits y consenso actuales.",
                "El test ya fue consultado en agosto y no debe describirse como completamente virgen.",
                "Ejecutar build_stage1_experiments_excel.py tras añadir resultados nuevos.",
            ],
        }
    )

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumen", index=False)
        details.to_excel(writer, sheet_name="Metricas_detalle", index=False)
        folds.to_excel(writer, sheet_name="CV_por_fold", index=False)
        current_config.to_excel(writer, sheet_name="Configuraciones", index=False)
        notes.to_excel(writer, sheet_name="Notas_y_protocolo", index=False)

    workbook = load_workbook(OUTPUT_PATH)
    for ws, name in [
        (workbook["Resumen"], "Stage1Summary"),
        (workbook["Metricas_detalle"], "Stage1Metrics"),
        (workbook["CV_por_fold"], "Stage1Folds"),
        (workbook["Configuraciones"], "Stage1Configs"),
        (workbook["Notas_y_protocolo"], "Stage1Notes"),
    ]:
        add_table(ws, name)
    workbook.save(OUTPUT_PATH)
    style_workbook()
    return OUTPUT_PATH


def historical_quality_metrics() -> tuple[dict[str, Any], pd.DataFrame, pd.DataFrame]:
    """Lee el experimento antiguo train-mejor/test-poor para documentarlo."""
    summary: dict[str, Any] = {
        "ID": "S1-H02",
        "Estado": "Historico",
        "Experimento": "MFCC117 + RF (audio-quality poor antiguo)",
        "Features": "13 MFCC + delta + delta2; mean/std/max = 117",
        "Clasificador": "Random Forest",
        "Datos/splits": "features_extracted_experiment_poor (10-08-2026)",
        "Umbral": 0.5,
        "Nota": (
            "Resultado historico train-mejor/test-peor. Tenia 325 grabaciones "
            "repartidas entre varios folds y TEST mezclaba calidades."
        ),
        "Fuente": str(HISTORICAL_QUALITY_TEXT),
    }
    if not HISTORICAL_QUALITY_TEXT.is_file():
        summary["Estado"] = "No encontrado"
        return summary, pd.DataFrame(), pd.DataFrame()

    text = HISTORICAL_QUALITY_TEXT.read_text(encoding="utf-8", errors="ignore")
    folds = [
        {"ID": "S1-H02", "fold": int(fold), "f1_cough": float(score)}
        for fold, score in re.findall(
            r"Fold\s+(\d+)\s+-\s+F1-Score:\s*([0-9.]+)", text
        )
    ]
    cv_match = re.search(r"VALIDACI[^:]*:\s*([0-9.]+)", text, flags=re.IGNORECASE)
    if cv_match:
        summary["F1 CV medio"] = float(cv_match.group(1))
    if folds:
        summary["F1 CV std"] = float(
            np.std([row["f1_cough"] for row in folds], ddof=0)
        )

    block_pattern = re.compile(
        r"Conjunto\s+(TRAIN|VALIDATION|TEST \(BLIND\))\s*:(.*?)(?=Conjunto|Pipeline|\Z)",
        flags=re.DOTALL,
    )
    patterns = {
        "accuracy": r"Accuracy:\s*([0-9.]+)",
        "precision_cough": r"Precision:\s*([0-9.]+)",
        "recall_cough": r"Recall:\s*([0-9.]+)",
        "f1_cough": r"F1-Score:\s*([0-9.]+)",
        "roc_auc": r"ROC-AUC:\s*([0-9.]+)",
    }
    split_names = {
        "TRAIN": "train_fit",
        "VALIDATION": "validation",
        "TEST (BLIND)": "test",
    }
    details: list[dict[str, Any]] = []
    for raw_split, block in block_pattern.findall(text):
        row: dict[str, Any] = {"ID": "S1-H02", "dataset": split_names[raw_split]}
        for name, pattern in patterns.items():
            match = re.search(pattern, block)
            row[name] = float(match.group(1)) if match else np.nan
        details.append(row)

    detail_df = pd.DataFrame(details)
    if not detail_df.empty:
        validation = detail_df[detail_df["dataset"] == "validation"]
        test = detail_df[detail_df["dataset"] == "test"]
        if not validation.empty:
            row = validation.iloc[0]
            summary.update(
                {
                    "F1 validation": row.get("f1_cough", np.nan),
                    "Recall tos validation": row.get("recall_cough", np.nan),
                    "Precision tos validation": row.get("precision_cough", np.nan),
                    "AUC validation": row.get("roc_auc", np.nan),
                }
            )
        if not test.empty:
            row = test.iloc[0]
            summary.update(
                {
                    "F1 test": row.get("f1_cough", np.nan),
                    "Recall tos test": row.get("recall_cough", np.nan),
                    "AUC test": row.get("roc_auc", np.nan),
                }
            )
    return summary, detail_df, pd.DataFrame(folds)


def quality_metrics() -> tuple[
    dict[str, Any], pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame
]:
    """Carga la repeticion corregida con las toses poor reservadas en TEST."""
    metrics_path = QUALITY_RESULTS / "metrics_summary.csv"
    summary: dict[str, Any] = {
        "ID": "S1-Q01",
        "Estado": "Pendiente de ejecucion",
        "Experimento": "MFCC117 + RF (train mejor / test poor)",
        "Features": "13 MFCC + delta + delta2; mean/std/max = 117",
        "Clasificador": "Random Forest",
        "Datos/splits": "multiclass_4c actual; todas las toses poor reservadas en TEST",
        "Umbral": 0.5,
        "Nota": (
            "Repeticion corregida: folds por grabacion y recall de tos "
            "desglosado por good/ok/poor."
        ),
        "Fuente": str(metrics_path),
    }
    if not metrics_path.is_file():
        return summary, pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    metrics = pd.read_csv(metrics_path)
    metrics.insert(0, "ID", "S1-Q01")
    folds_path = QUALITY_RESULTS / "cv_fold_metrics.csv"
    folds = pd.read_csv(folds_path) if folds_path.is_file() else pd.DataFrame()
    if not folds.empty:
        folds.insert(0, "ID", "S1-Q01")
    config_path = QUALITY_RESULTS / "experiment_configuration.csv"
    config = pd.read_csv(config_path) if config_path.is_file() else pd.DataFrame()
    if not config.empty:
        config.insert(0, "ID", "S1-Q01")
    quality_path = QUALITY_RESULTS / "test_cough_recall_by_quality.csv"
    quality = pd.read_csv(quality_path) if quality_path.is_file() else pd.DataFrame()
    if not quality.empty:
        quality.insert(0, "ID", "S1-Q01")

    summary["Estado"] = "Finalizado"
    oof = metrics[metrics["dataset"] == "train_oof"]
    validation = metrics[metrics["dataset"] == "validation"]
    test = metrics[metrics["dataset"] == "test"]
    if not oof.empty:
        summary["Macro-F1 OOF"] = oof.iloc[0].get("macro_f1", np.nan)
    if not folds.empty:
        summary["F1 CV medio"] = folds["f1_cough"].mean()
        summary["F1 CV std"] = folds["f1_cough"].std(ddof=0)
    if not validation.empty:
        row = validation.iloc[0]
        summary.update(
            {
                "F1 validation": row.get("f1_cough", np.nan),
                "Macro-F1 validation": row.get("macro_f1", np.nan),
                "Recall tos validation": row.get("recall_cough", np.nan),
                "Especificidad validation": row.get("specificity_no_cough", np.nan),
                "Precision tos validation": row.get("precision_cough", np.nan),
                "AUC validation": row.get("roc_auc", np.nan),
            }
        )
    if not test.empty:
        row = test.iloc[0]
        summary.update(
            {
                "F1 test": row.get("f1_cough", np.nan),
                "Macro-F1 test": row.get("macro_f1", np.nan),
                "Recall tos test": row.get("recall_cough", np.nan),
                "Especificidad test": row.get("specificity_no_cough", np.nan),
                "Precision tos test": row.get("precision_cough", np.nan),
                "AUC test": row.get("roc_auc", np.nan),
            }
        )
    if not quality.empty:
        for level in ["good", "ok", "poor"]:
            selected = quality[quality["cough_quality"] == level]
            if not selected.empty:
                summary[f"Recall tos test {level}"] = selected.iloc[0].get(
                    "recall_cough", np.nan
                )
    if not config.empty:
        config_map = dict(zip(config["parameter"], config["value"]))
        summary["Configuracion ganadora"] = (
            f"RF trees={config_map.get('n_estimators', '?')}, "
            f"depth={config_map.get('max_depth', '?')}, "
            f"leaf={config_map.get('min_samples_leaf', '?')}"
        )
        if "model_size_kb" in config_map:
            summary["Tamano modelo KB"] = config_map["model_size_kb"]
        if "tree_nodes_total" in config_map:
            summary["Nodos arboles"] = config_map["tree_nodes_total"]
    return summary, metrics, folds, config, quality


def compact_metrics(
    experiment_id: str,
    results_dir: Path,
    experiment_name: str,
    classifier_name: str,
) -> tuple[
    dict[str, Any],
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
]:
    """Carga un ganador compacto y sus candidatos de busqueda OOF."""
    metrics_path = results_dir / "metrics_summary.csv"
    random_protocol = results_dir.parent.name.endswith("_random")
    summary: dict[str, Any] = {
        "ID": experiment_id,
        "Estado": "Pendiente de ejecucion",
        "Experimento": experiment_name,
        "Features": "13 MFCC + delta + delta2; mean/std/max = 117",
        "Clasificador": classifier_name,
        "Datos/splits": (
            "splits random actuales; ninguna calidad forzada a TEST"
            if random_protocol
            else "mismos splits audio-quality poor de S1-Q01"
        ),
        "Umbral": 0.5,
        "Nota": "Hiperparametros elegidos solo mediante OOF de TRAIN.",
        "Fuente": str(metrics_path),
    }
    if not metrics_path.is_file():
        empty = pd.DataFrame()
        return summary, empty, empty, empty, empty, empty

    metrics = pd.read_csv(metrics_path)
    metrics.insert(0, "ID", experiment_id)
    folds_path = results_dir / "best_cv_fold_metrics.csv"
    folds = pd.read_csv(folds_path) if folds_path.is_file() else pd.DataFrame()
    if not folds.empty:
        folds.insert(0, "ID", experiment_id)
    config_path = results_dir / "experiment_configuration.csv"
    config = pd.read_csv(config_path) if config_path.is_file() else pd.DataFrame()
    if not config.empty:
        config.insert(0, "ID", experiment_id)
    quality_path = results_dir / "test_cough_recall_by_quality.csv"
    quality = pd.read_csv(quality_path) if quality_path.is_file() else pd.DataFrame()
    if not quality.empty:
        quality.insert(0, "ID", experiment_id)
    candidates_path = results_dir / "candidate_cv_results.csv"
    candidates = (
        pd.read_csv(candidates_path) if candidates_path.is_file() else pd.DataFrame()
    )
    if not candidates.empty:
        candidates.insert(0, "ID", experiment_id)

    summary["Estado"] = "Finalizado"
    oof = metrics[metrics["dataset"] == "train_oof"]
    validation = metrics[metrics["dataset"] == "validation"]
    test = metrics[metrics["dataset"] == "test"]
    if not oof.empty:
        row = oof.iloc[0]
        summary["Macro-F1 OOF"] = row.get("macro_f1", np.nan)
        summary["Recall tos OOF"] = row.get("recall_cough", np.nan)
    if not folds.empty:
        summary["F1 CV medio"] = folds["f1_cough"].mean()
        summary["F1 CV std"] = folds["f1_cough"].std(ddof=0)
    if not validation.empty:
        row = validation.iloc[0]
        summary.update(
            {
                "F1 validation": row.get("f1_cough", np.nan),
                "Macro-F1 validation": row.get("macro_f1", np.nan),
                "Recall tos validation": row.get("recall_cough", np.nan),
                "Especificidad validation": row.get("specificity_no_cough", np.nan),
                "Precision tos validation": row.get("precision_cough", np.nan),
                "AUC validation": row.get("roc_auc", np.nan),
            }
        )
    if not test.empty:
        row = test.iloc[0]
        summary.update(
            {
                "F1 test": row.get("f1_cough", np.nan),
                "Macro-F1 test": row.get("macro_f1", np.nan),
                "Recall tos test": row.get("recall_cough", np.nan),
                "Especificidad test": row.get("specificity_no_cough", np.nan),
                "Precision tos test": row.get("precision_cough", np.nan),
                "AUC test": row.get("roc_auc", np.nan),
            }
        )
    if not quality.empty:
        for level in ["good", "ok", "poor"]:
            selected = quality[quality["cough_quality"] == level]
            if not selected.empty:
                summary[f"Recall tos test {level}"] = selected.iloc[0].get(
                    "recall_cough", np.nan
                )
    if not config.empty:
        config_map = dict(zip(config["parameter"], config["value"]))
        for source, destination in [
            ("winner", "Configuracion ganadora"),
            ("model_size_kb", "Tamano modelo KB"),
            ("desktop_single_sample_ms", "Inferencia escritorio ms"),
            ("trainable_parameters", "Parametros entrenables"),
            ("deployment_float_values", "Valores float despliegue"),
            ("tree_nodes_total", "Nodos arboles"),
        ]:
            if source in config_map:
                summary[destination] = config_map[source]
    return summary, metrics, folds, config, quality, candidates


def wst_o01_recording_metrics() -> tuple[
    dict[str, Any],
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
]:
    """Carga el experimento WST S0+S1 o informa su fase actual."""

    experiment_id = "S1-W01"
    metrics_path = WST_O01_RECORDING_LR_RESULTS / "metrics_summary.csv"
    phase_a_path = (
        WST_O01_RECORDING_LR_RESULTS / "phase_a_candidate_cv_results.csv"
    )
    phase_b_path = (
        WST_O01_RECORDING_LR_RESULTS / "phase_b_candidate_cv_results.csv"
    )
    summary: dict[str, Any] = {
        "ID": experiment_id,
        "Estado": "Pendiente de ejecucion",
        "Experimento": "WST S0+S1 recording mejorado + LR",
        "Features": (
            "95 paths S0+S1 por ventana; raw/log; weighted mean/std + max; "
            "285/291/294 por grabacion"
        ),
        "Clasificador": "Logistic Regression",
        "Datos/splits": (
            "splits random agrupados con toses FSD50K revisadas; una fila "
            "y etiqueta por original_uuid"
        ),
        "Umbral": "Seleccionado con OOF de TRAIN",
        "Nota": (
            "Compara raw/log-WST, amplitud, PCA, C, L1/L2 y class_weight. "
            "Registra por separado recall de tos COUGHVID y de las nuevas "
            "toses FSD50K para auditar el cambio de dominio."
        ),
        "Fuente": str(metrics_path),
    }

    if not metrics_path.is_file():
        if phase_b_path.is_file():
            summary["Estado"] = "En ejecucion (Fase B)"
        elif phase_a_path.is_file():
            summary["Estado"] = "En ejecucion (Fase A cerrada)"
        empty = pd.DataFrame()
        candidates = []
        for phase, path in (("A", phase_a_path), ("B", phase_b_path)):
            if not path.is_file():
                continue
            frame = pd.read_csv(path)
            frame.insert(0, "search_phase", phase)
            frame.insert(0, "ID", experiment_id)
            candidates.append(frame)
        candidate_frame = (
            pd.concat(candidates, ignore_index=True, sort=False)
            if candidates
            else empty
        )
        return summary, empty, empty, empty, candidate_frame

    metrics = pd.read_csv(metrics_path)
    metrics.insert(0, "ID", experiment_id)
    folds_path = WST_O01_RECORDING_LR_RESULTS / "best_cv_fold_metrics.csv"
    folds = pd.read_csv(folds_path) if folds_path.is_file() else pd.DataFrame()
    if not folds.empty:
        folds.insert(0, "ID", experiment_id)
    config_path = WST_O01_RECORDING_LR_RESULTS / "experiment_configuration.csv"
    config = pd.read_csv(config_path) if config_path.is_file() else pd.DataFrame()
    if not config.empty:
        config.insert(0, "ID", experiment_id)

    candidate_frames = []
    for phase, path in (("A", phase_a_path), ("B", phase_b_path)):
        if not path.is_file():
            continue
        frame = pd.read_csv(path)
        frame.insert(0, "search_phase", phase)
        frame.insert(0, "ID", experiment_id)
        candidate_frames.append(frame)
    candidates = (
        pd.concat(candidate_frames, ignore_index=True, sort=False)
        if candidate_frames
        else pd.DataFrame()
    )

    summary["Estado"] = "Finalizado"
    oof = metrics[metrics["dataset"].astype(str).eq("train_oof")]
    validation = metrics[metrics["dataset"].astype(str).eq("validation")]
    test = metrics[metrics["dataset"].astype(str).eq("test")]
    if not oof.empty:
        row = oof.iloc[0]
        summary.update(
            {
                "Macro-F1 OOF": row.get("macro_f1", np.nan),
                "Recall tos OOF": row.get("recall_cough", np.nan),
                "Recall tos original OOF": row.get(
                    "original_cough_recall", np.nan
                ),
                "Recall tos FSD50K OOF": row.get(
                    "new_fsd50k_cough_recall", np.nan
                ),
            }
        )
    if not folds.empty and "f1_cough" in folds:
        summary["F1 CV medio"] = folds["f1_cough"].mean()
        summary["F1 CV std"] = folds["f1_cough"].std(ddof=0)
    if not validation.empty:
        row = validation.iloc[0]
        summary.update(
            {
                "F1 validation": row.get("f1_cough", np.nan),
                "Macro-F1 validation": row.get("macro_f1", np.nan),
                "Recall tos validation": row.get("recall_cough", np.nan),
                "Especificidad validation": row.get(
                    "specificity_no_cough", np.nan
                ),
                "Precision tos validation": row.get(
                    "precision_cough", np.nan
                ),
                "AUC validation": row.get("roc_auc", np.nan),
                "Recall tos original validation": row.get(
                    "original_cough_recall", np.nan
                ),
                "Recall tos FSD50K validation": row.get(
                    "new_fsd50k_cough_recall", np.nan
                ),
            }
        )
    if not test.empty:
        row = test.iloc[0]
        summary.update(
            {
                "F1 test": row.get("f1_cough", np.nan),
                "Macro-F1 test": row.get("macro_f1", np.nan),
                "Recall tos test": row.get("recall_cough", np.nan),
                "Especificidad test": row.get("specificity_no_cough", np.nan),
                "Precision tos test": row.get("precision_cough", np.nan),
                "AUC test": row.get("roc_auc", np.nan),
                "Recall tos original test": row.get(
                    "original_cough_recall", np.nan
                ),
                "Recall tos FSD50K test": row.get(
                    "new_fsd50k_cough_recall", np.nan
                ),
            }
        )
    if not config.empty and {"parameter", "value"}.issubset(config.columns):
        config_map = dict(zip(config["parameter"], config["value"]))
        for source, destination in (
            ("winner", "Configuracion ganadora"),
            ("model_size_kb", "Tamano modelo KB"),
            ("desktop_classifier_inference_ms", "Inferencia escritorio ms"),
        ):
            if source in config_map:
                summary[destination] = config_map[source]
    return summary, metrics, folds, config, candidates


# Redefinicion intencionada: amplia el constructor original con los dos
# experimentos de calidad sin alterar los CSV historicos.
def build_workbook() -> Path:
    historical_summary, historical_detail, historical_folds = historical_metrics()
    quality_h_summary, quality_h_detail, quality_h_folds = historical_quality_metrics()
    current_summary, current_detail, current_folds, current_config = current_metrics()
    quality_summary, quality_detail, quality_folds, quality_config, quality_by_level = (
        quality_metrics()
    )
    compact_lr = compact_metrics(
        "S1-Q02",
        COMPACT_LR_RESULTS,
        "MFCC117 + LR compacta (train mejor / test poor)",
        "Logistic Regression",
    )
    compact_rf = compact_metrics(
        "S1-Q03",
        COMPACT_RF_RESULTS,
        "MFCC117 + RF compacto (train mejor / test poor)",
        "Random Forest compacto",
    )
    random_lr = compact_metrics(
        "S1-R02",
        RANDOM_LR_RESULTS,
        "MFCC117 + LR compacta (splits random)",
        "Logistic Regression",
    )
    random_rf = compact_metrics(
        "S1-R03",
        RANDOM_RF_RESULTS,
        "MFCC117 + RF compacto (splits random)",
        "Random Forest compacto",
    )
    fsd50k_cough_lr = compact_metrics(
        "S1-R04",
        FSD50K_COUGH_LR_RESULTS,
        "MFCC117 + LR fija + toses FSD50K revisadas",
        "Logistic Regression",
    )
    fsd50k_cough_lr[0]["Datos/splits"] = (
        "splits random agrupados; incorpora 172 toses FSD50K revisadas "
        "manualmente y excluidas de Stage 2"
    )
    fsd50k_cough_lr[0]["Nota"] = (
        "Configuracion LR congelada del experimento S1-R02; sin nueva busqueda "
        "de hiperparametros ni umbral."
    )
    fsd50k_cough_lr_search = compact_metrics(
        "S1-R05",
        FSD50K_COUGH_LR_SEARCH_RESULTS,
        "MFCC117 + LR optimizada + toses FSD50K revisadas",
        "Logistic Regression",
    )
    fsd50k_cough_lr_search[0]["Datos/splits"] = (
        "mismos splits agrupados de S1-R04; 172 toses FSD50K revisadas y "
        "excluidas de Stage 2"
    )
    fsd50k_cough_lr_search[0]["Nota"] = (
        "Busqueda OOF en dos fases; seleccion sensible al rendimiento dentro "
        "de FSD50K y con TEST reservado."
    )
    wst_o01 = wst_o01_recording_metrics()

    summary = pd.DataFrame(
        [
            historical_summary,
            quality_h_summary,
            current_summary,
            quality_summary,
            compact_lr[0],
            compact_rf[0],
            random_lr[0],
            random_rf[0],
            fsd50k_cough_lr[0],
            fsd50k_cough_lr_search[0],
            wst_o01[0],
        ]
    )
    preferred_columns = [
        "ID", "Estado", "Experimento", "Features", "Clasificador", "Datos/splits",
        "Umbral", "F1 CV medio", "F1 CV std", "Macro-F1 OOF", "Recall tos OOF",
        "Recall tos original OOF", "Recall tos FSD50K OOF",
        "F1 validation",
        "Macro-F1 validation", "Recall tos validation", "Especificidad validation",
        "Precision tos validation", "AUC validation", "F1 test", "Macro-F1 test",
        "Recall tos original validation", "Recall tos FSD50K validation",
        "Recall tos test", "Especificidad test", "Precision tos test", "AUC test",
        "Recall tos original test", "Recall tos FSD50K test",
        "Recall tos test good", "Recall tos test ok", "Recall tos test poor",
        "Configuracion ganadora", "Tamano modelo KB", "Inferencia escritorio ms",
        "Parametros entrenables", "Valores float despliegue", "Nodos arboles",
        "Nota", "Fuente",
    ]
    summary = summary.reindex(columns=preferred_columns)
    details = pd.concat(
        [
            historical_detail,
            quality_h_detail,
            current_detail,
            quality_detail,
            compact_lr[1],
            compact_rf[1],
            random_lr[1],
            random_rf[1],
            fsd50k_cough_lr[1],
            fsd50k_cough_lr_search[1],
            wst_o01[1],
        ],
        ignore_index=True,
        sort=False,
    )
    folds = pd.concat(
        [
            historical_folds,
            quality_h_folds,
            current_folds,
            quality_folds,
            compact_lr[2],
            compact_rf[2],
            random_lr[2],
            random_rf[2],
            fsd50k_cough_lr[2],
            fsd50k_cough_lr_search[2],
            wst_o01[2],
        ],
        ignore_index=True,
        sort=False,
    )
    configurations = pd.concat(
        [
            current_config,
            quality_config,
            compact_lr[3],
            compact_rf[3],
            random_lr[3],
            random_rf[3],
            fsd50k_cough_lr[3],
            fsd50k_cough_lr_search[3],
            wst_o01[3],
        ],
        ignore_index=True,
        sort=False,
    )
    quality_all = pd.concat(
        [
            quality_by_level,
            compact_lr[4],
            compact_rf[4],
            random_lr[4],
            random_rf[4],
            fsd50k_cough_lr[4],
            fsd50k_cough_lr_search[4],
        ],
        ignore_index=True,
        sort=False,
    )
    compact_candidates = pd.concat(
        [
            compact_lr[5],
            compact_rf[5],
            random_lr[5],
            random_rf[5],
            fsd50k_cough_lr[5],
            fsd50k_cough_lr_search[5],
            wst_o01[4],
        ],
        ignore_index=True,
        sort=False,
    )
    notes = pd.DataFrame(
        {
            "Aspecto": [
                "Objetivo", "Mapeo", "Comparabilidad", "Protocolo quality",
                "TEST mixto", "Control de fuga", "Modelos compactos",
                "Comparacion random", "Cambio de dominio WST", "Actualizacion",
            ],
            "Descripcion": [
                "Stage 1 detecta tos frente a no tos.",
                "Clase 0=no_cough; dry, wet y unknown se consideran tos (clase 1).",
                "Los historicos usan metadatos antiguos; S1-C01 y S1-Q01 usan el consenso actual.",
                "S1-Q01 excluye de TRAIN y validation todas las toses de calidad poor.",
                (
                    "TEST contiene todas las poor y una muestra adicional de otras "
                    "calidades/no-tos; por eso se informa recall por calidad."
                ),
                (
                    "Los folds nuevos se asignan por original_uuid antes de segmentar; "
                    "los segmentos de una grabacion nunca cruzan folds."
                ),
                (
                    "S1-Q02 y S1-Q03 usan los mismos datos que S1-Q01; seleccionan "
                    "hiperparametros solo mediante OOF y registran tamano y latencia de escritorio."
                ),
                (
                    "S1-R02 y S1-R03 repiten LR y RF compacto sobre los splits random "
                    "actuales, sin forzar ninguna calidad de tos a TEST."
                ),
                (
                    "S1-W01 informa recall de tos original y FSD50K por separado; "
                    "el origen esta fuertemente correlacionado con la clase y la "
                    "metrica global puede ocultar una mala generalizacion FSD50K."
                ),
                "Los entrenamientos actuales actualizan este Excel automaticamente.",
            ],
        }
    )

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumen", index=False)
        details.to_excel(writer, sheet_name="Metricas_detalle", index=False)
        folds.to_excel(writer, sheet_name="CV_por_fold", index=False)
        configurations.to_excel(writer, sheet_name="Configuraciones", index=False)
        quality_all.to_excel(writer, sheet_name="Recall_por_calidad", index=False)
        compact_candidates.to_excel(
            writer, sheet_name="Candidatos_compactos", index=False
        )
        notes.to_excel(writer, sheet_name="Notas_y_protocolo", index=False)

    workbook = load_workbook(OUTPUT_PATH)
    for ws, name in [
        (workbook["Resumen"], "Stage1Summary"),
        (workbook["Metricas_detalle"], "Stage1Metrics"),
        (workbook["CV_por_fold"], "Stage1Folds"),
        (workbook["Configuraciones"], "Stage1Configs"),
        (workbook["Recall_por_calidad"], "Stage1Quality"),
        (workbook["Candidatos_compactos"], "Stage1CompactCandidates"),
        (workbook["Notas_y_protocolo"], "Stage1Notes"),
    ]:
        add_table(ws, name)
    workbook.save(OUTPUT_PATH)
    style_workbook()
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_workbook()
    print(f"Excel Stage 1 generado: {path}")
